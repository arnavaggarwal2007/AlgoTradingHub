"""
================================================================================
UNIT TESTS: INSTITUTIONAL WEEKLY SWING TRADING SCREENER & WATCHLIST GENERATOR
================================================================================
Tests:
1. Universe constituent fetching (Dow 30, Nasdaq 100, S&P 500, fallback)
2. Indicator calculations (EMA10, EMA21, SMA50, SMA150, SMA200, ATR14, slopes)
3. 30-Day Swing Structure Analyzer (Higher Lows SL2 >= SL1, Volatility Contraction)
4. TradingView 13-Rules Evaluation & Gating Logic
5. Institutional Quality Score (IQS: 0-100) and Star Ratings
6. Automated Trade Plan Parameters (Stop loss bounds, profit targets, R:R)
7. Multi-Sheet Excel Workbook Export
================================================================================
"""

import os
import tempfile
import pytest
import numpy as np
import pandas as pd
import openpyxl

from scripts.scanner.run_institutional_weekly_screener import (
    get_screener_universe,
    calculate_screener_indicators,
    analyze_swing_structure,
    evaluate_screener_candidate,
    export_screener_to_excel,
    DOW_30_SYMBOLS,
    NASDAQ_100_SYMBOLS
)


def create_synthetic_market_data(
    bars: int = 260,
    base_price: float = 100.0,
    trend: str = "uptrend",
    vol_base: float = 500000.0
) -> pd.DataFrame:
    """Helper to generate clean synthetic OHLCV daily bars"""
    dates = pd.date_range(end=pd.Timestamp.today(), periods=bars, freq='B')
    
    if trend == "uptrend":
        # Smooth steady upward progression
        price_trend = np.linspace(base_price * 0.70, base_price, bars)
    elif trend == "downtrend":
        price_trend = np.linspace(base_price * 1.30, base_price, bars)
    else:
        price_trend = np.full(bars, base_price)
        
    # Add minor noise
    noise = np.sin(np.linspace(0, 10, bars)) * (base_price * 0.02)
    closes = price_trend + noise
    highs = closes * 1.015
    lows = closes * 0.985
    opens = (closes + lows) / 2.0
    volumes = np.full(bars, vol_base)
    # Ensure last bar has volume increase
    volumes[-1] = vol_base * 1.30
    
    df = pd.DataFrame({
        'Open': opens,
        'High': highs,
        'Low': lows,
        'Close': closes,
        'Volume': volumes
    }, index=dates)
    return df


class TestInstitutionalScreener:
    
    def test_get_screener_universe_dow30(self):
        """Dow 30 constituents are properly retrieved and deduplicated"""
        dow = get_screener_universe("dow30")
        assert len(dow) >= 30
        assert "AAPL" in dow
        assert "MSFT" in dow
        assert "NVDA" in dow
        assert dow == sorted(dow)
        
    def test_get_screener_universe_nasdaq100(self):
        """Nasdaq 100 constituents are properly retrieved"""
        nasdaq = get_screener_universe("nasdaq100")
        assert len(nasdaq) >= 90
        assert "AAPL" in nasdaq
        assert "AMZN" in nasdaq
        assert "GOOGL" in nasdaq
        
    def test_get_screener_universe_combined(self):
        """Combined universe merges indices without duplicates"""
        combined = get_screener_universe("combined")
        assert len(combined) > len(DOW_30_SYMBOLS)
        assert len(combined) == len(set(combined))

    def test_calculate_screener_indicators_insufficient_data(self):
        """Returns None if dataframe has fewer than 200 bars"""
        df_short = create_synthetic_market_data(bars=150)
        res = calculate_screener_indicators(df_short)
        assert res is None

    def test_calculate_screener_indicators_success(self):
        """Correctly computes all required technical indicators and slopes"""
        df = create_synthetic_market_data(bars=260, base_price=150.0, trend="uptrend")
        res = calculate_screener_indicators(df)
        assert res is not None
        
        expected_cols = [
            'EMA10', 'EMA21', 'SMA50', 'SMA150', 'SMA200',
            'VOL_SMA21', 'VOL_SMA100', 'HIGH_52W', 'LOW_52W',
            'ATR14', 'SLOPE_EMA21', 'SLOPE_SMA50', 'SLOPE_SMA200'
        ]
        for col in expected_cols:
            assert col in res.columns
            assert not pd.isna(res[col].iloc[-1])
            
        # In a steady uptrend, moving averages should align bullishly
        last = res.iloc[-1]
        assert last['EMA10'] > last['EMA21']
        assert last['EMA21'] > last['SMA50']
        assert last['SMA50'] > last['SMA200']
        assert last['SLOPE_EMA21'] > 0
        assert last['SLOPE_SMA50'] > 0

    def test_analyze_swing_structure_higher_low_and_contraction(self):
        """Verifies detection of 2 swing highs, 2 swing lows with SL2 >= SL1 and contraction"""
        dates = pd.date_range(end=pd.Timestamp.today(), periods=35, freq='B')
        df = pd.DataFrame(index=dates)
        
        # Construct synthetic swings:
        # Day 5: Low 100 (SL1)
        # Day 12: High 120 (SH1) -> Range1 = 20
        # Day 20: Low 105 (SL2 > SL1)
        # Day 27: High 118 (SH2) -> Range2 = 13 (Contraction)
        # Day 34: Current close
        prices = np.full(35, 110.0)
        # SL1 at index 5
        prices[4] = 104.0; prices[5] = 100.0; prices[6] = 106.0
        # SH1 at index 12
        prices[11] = 115.0; prices[12] = 120.0; prices[13] = 114.0
        # SL2 at index 20
        prices[19] = 108.0; prices[20] = 105.0; prices[21] = 109.0
        # SH2 at index 27
        prices[26] = 114.0; prices[27] = 118.0; prices[28] = 113.0
        
        df['High'] = prices * 1.01
        df['Low'] = prices * 0.99
        # Specifically anchor peaks and troughs
        df.loc[dates[5], 'Low'] = 100.0
        df.loc[dates[12], 'High'] = 120.0
        df.loc[dates[20], 'Low'] = 105.0
        df.loc[dates[27], 'High'] = 118.0
        df['Close'] = prices
        df['Open'] = prices
        df['Volume'] = 500000.0
        
        swing_res = analyze_swing_structure(df, window_bars=30)
        assert swing_res['valid_structure'] is True
        assert swing_res['swing_highs_count'] >= 2
        assert swing_res['swing_lows_count'] >= 2
        assert swing_res['higher_low'] is True
        assert swing_res['contraction'] is True

    def test_evaluate_screener_candidate_passes_all_13_rules(self):
        """A textbook Stage 2 leader passes all 13 TradingView criteria and receives trade plan"""
        df = create_synthetic_market_data(bars=260, base_price=100.0, trend="uptrend", vol_base=600000.0)
        df_ind = calculate_screener_indicators(df)
        
        eval_res = evaluate_screener_candidate("LEAD", df_ind)
        assert eval_res is not None
        assert eval_res['symbol'] == "LEAD"
        assert eval_res['close'] > 10.0
        assert eval_res['vol_100d_k'] >= 250.0
        assert eval_res['ytd_perf_pct'] > 8.0
        assert eval_res['iqs_score'] >= 60
        assert "⭐" in eval_res['star_rating']
        
        # Check automated trade plan calculations
        entry = eval_res['suggested_entry']
        stop = eval_res['initial_stop_loss']
        assert stop < entry
        assert 5.0 <= eval_res['risk_pct'] <= 10.0
        assert eval_res['target_1'] > entry
        assert eval_res['target_2'] > eval_res['target_1']
        assert eval_res['target_3'] > eval_res['target_2']
        assert eval_res['risk_reward'] >= 1.0

    def test_evaluate_screener_candidate_fails_price_under_10(self):
        """Fails Rule 1 if price < 10 USD"""
        df = create_synthetic_market_data(bars=260, base_price=8.0, trend="uptrend")
        df_ind = calculate_screener_indicators(df)
        assert evaluate_screener_candidate("PENNY", df_ind) is None

    def test_evaluate_screener_candidate_fails_low_volume(self):
        """Fails Rule 11 if 100-day average volume < 250,000 shares"""
        df = create_synthetic_market_data(bars=260, base_price=100.0, vol_base=100000.0)
        df_ind = calculate_screener_indicators(df)
        assert evaluate_screener_candidate("ILLIQ", df_ind) is None

    def test_evaluate_screener_candidate_fails_bearish_trend(self):
        """Fails Rule 4, 8, 10, 12, 13 if moving averages are in downtrend"""
        df = create_synthetic_market_data(bars=260, base_price=100.0, trend="downtrend")
        df_ind = calculate_screener_indicators(df)
        assert evaluate_screener_candidate("BEAR", df_ind) is None

    def test_evaluate_screener_candidate_fails_over_extended_from_sma50(self):
        """Fails Rule 6 if price is > 10% above SMA50"""
        df = create_synthetic_market_data(bars=260, base_price=100.0, trend="uptrend")
        df_ind = calculate_screener_indicators(df)
        # Artificially spike the final close to 30% above SMA50
        df_ind.loc[df_ind.index[-1], 'Close'] = float(df_ind['SMA50'].iloc[-1]) * 1.30
        assert evaluate_screener_candidate("SPIKE", df_ind) is None

    def test_export_screener_to_excel(self):
        """Excel exporter creates a multi-sheet formatted workbook with expected headers"""
        # Generate 3 mock candidates
        mock_candidates = [
            {
                'symbol': 'NVDA', 'close': 140.50, 'ytd_perf_pct': 185.2, 'vol_100d_k': 45200.0,
                'rvol': 1.65, 'dist_to_ema21_pct': 1.2, 'dist_to_sma50_pct': 3.1, 'dist_to_sma200_pct': 28.5,
                'dist_from_52w_high_pct': 2.5, 'dist_from_52w_low_pct': 110.0, 'slope_ema21': 3.2,
                'slope_sma50': 2.1, 'ema21_slope_desc': 'Strongly Upward', 'sma50_slope_desc': 'Strongly Upward',
                'buy_zone_type': 'Dual Convergence Sweet Spot', 'swing_higher_low': 'YES (SL2 >= SL1)',
                'swing_contraction': 'YES (VCP)', 'sl1': 130.0, 'sl2': 136.0, 'sh1': 142.0, 'sh2': 144.0,
                'rs_score': 94.5, 'iqs_score': 96, 'star_rating': '⭐⭐⭐⭐⭐ (Grade A+ Elite)',
                'action_verdict': 'PRIME BUY SETUP', 'suggested_entry': 140.50, 'initial_stop_loss': 132.00,
                'risk_pct': 6.0, 'target_1': 147.53, 'target_2': 154.55, 'target_3': 161.58, 'risk_reward': 1.65
            },
            {
                'symbol': 'MSFT', 'close': 425.00, 'ytd_perf_pct': 18.5, 'vol_100d_k': 18500.0,
                'rvol': 1.25, 'dist_to_ema21_pct': 2.1, 'dist_to_sma50_pct': 2.8, 'dist_to_sma200_pct': 14.2,
                'dist_from_52w_high_pct': 5.2, 'dist_from_52w_low_pct': 35.0, 'slope_ema21': 1.2,
                'slope_sma50': 0.8, 'ema21_slope_desc': 'Upward', 'sma50_slope_desc': 'Upward',
                'buy_zone_type': '21 EMA Prime Zone', 'swing_higher_low': 'YES (SL2 >= SL1)',
                'swing_contraction': 'YES (VCP)', 'sl1': 410.0, 'sl2': 418.0, 'sh1': 430.0, 'sh2': 432.0,
                'rs_score': 78.2, 'iqs_score': 88, 'star_rating': '⭐⭐⭐⭐ (Grade A High Conviction)',
                'action_verdict': 'HIGH WATCHLIST PRIORITY', 'suggested_entry': 425.00, 'initial_stop_loss': 399.50,
                'risk_pct': 6.0, 'target_1': 446.25, 'target_2': 467.50, 'target_3': 488.75, 'risk_reward': 1.67
            }
        ]
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            temp_path = tf.name
            
        try:
            export_screener_to_excel(mock_candidates, output_path=temp_path)
            assert os.path.exists(temp_path)
            
            wb = openpyxl.load_workbook(temp_path)
            assert "Top_Watchlist_Next_Week" in wb.sheetnames
            assert "All_Screened_Candidates" in wb.sheetnames
            assert "Watchlist_TXT_Export" in wb.sheetnames
            
            # Check sheet 1 contents
            ws_top = wb["Top_Watchlist_Next_Week"]
            assert ws_top.cell(row=1, column=1).value == "Rank"
            assert ws_top.cell(row=2, column=2).value == "NVDA"
            assert ws_top.cell(row=3, column=2).value == "MSFT"
            
            # Check sheet 3 contents
            ws_txt = wb["Watchlist_TXT_Export"]
            assert ws_txt.cell(row=1, column=1).value == "Symbol (Ready to Copy to config/watchlist.txt)"
            assert ws_txt.cell(row=2, column=1).value == "NVDA"
            assert ws_txt.cell(row=3, column=1).value == "MSFT"
            
            wb.close()
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
